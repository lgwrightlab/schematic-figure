"""Convert the SVG of a figure into a PowerPoint deck of NATIVE objects.

Every <text> becomes one text box with live text. Rectangles, ellipses and lines become preset shapes. Polygons and
paths become freeform shapes (custGeom). SVG groups become PowerPoint groups. Fills, strokes, dashes, opacity,
linear and radial gradients and drop shadows become PowerPoint formats. Any person in any app can click a label,
retype it, move a shape or recolour it.

python tools/svg2pptx_native.py out.pptx fig.svg [fig2.svg ...]
python tools/svg2pptx.py --native out.pptx fig.svg fig.png [fig2.svg fig2.png ...]    (same result, PNG not read)

The figure is fitted inside the 16:9 slide (338.7 x 190.5 mm), centred and never stretched. A 16:9 figure fills it.
One SVG user unit is one point, as the engine draws. The converter reads the SVG elements that the engine writes:
g, rect, circle, ellipse, line, polyline, polygon, path, text, tspan, textPath, image and the defs for gradients,
filters, patterns and clip paths. An unknown element is skipped and counted. It never fails the conversion.

What PowerPoint cannot hold is dropped or approximated, and the script prints a count for each kind:
  clip paths, masks, glow filters, patterns (solid fallback), tspan dx shifts, text halos (paint-order stroke),
  text on a curved path (straight text along the chord), group opacity (pushed down to each child).
"""
import base64
import copy
import io
import math
import os
import re
import sys
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

from lxml import etree

HERE = Path(__file__).resolve().parent
SVGNS = 'http://www.w3.org/2000/svg'
XLINK = 'http://www.w3.org/1999/xlink'
NS_DECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
SLIDE_W_MM, SLIDE_H_MM = 338.7, 190.5
EMU_MM = 36000
EMU_PT = 12700
MAX_GROUP_DEPTH = 3          # a deeper group is flattened into its parent
RADIAL_SCALE = 0.7071        # PowerPoint's circular gradient ends at the box corner, SVG's at radius r
# Text calibration, measured in PowerPoint for Mac: with 100 % line spacing the first baseline lies 0.9766 sizes
# below the box top, and a change of the spacing by p moves it by 0.3825 sizes per unit of p.
BASE_ASC = 0.9766
ASC_PER_P = 0.3825
SINGLE_LINE = 1.2            # PowerPoint's single line pitch, in font sizes
TEXT_NUDGE_EMU = 4400        # small text sat about 0.35 pt too high in PowerPoint, measured on the demos
SCRIPT_AUTO = 0.66          # PowerPoint draws a run with a baseline shift at this fraction of its sz (to measure)

INHERITED = ('fill', 'fill-opacity', 'stroke', 'stroke-opacity', 'stroke-width', 'stroke-linecap', 'stroke-linejoin',
             'stroke-dasharray', 'stroke-miterlimit', 'fill-rule', 'font-family', 'font-size', 'font-weight',
             'font-style', 'text-anchor', 'paint-order', 'letter-spacing', 'marker-start', 'marker-end',
             'text-decoration')
SKIP_TAGS = {'defs', 'clipPath', 'mask', 'marker', 'symbol', 'pattern', 'style', 'title', 'desc', 'metadata',
             'linearGradient', 'radialGradient', 'filter', 'script'}

NAMED = {'black': '000000', 'white': 'FFFFFF', 'red': 'FF0000', 'green': '008000', 'blue': '0000FF',
         'yellow': 'FFFF00', 'gray': '808080', 'grey': '808080', 'orange': 'FFA500', 'purple': '800080',
         'cyan': '00FFFF', 'magenta': 'FF00FF', 'silver': 'C0C0C0', 'navy': '000080', 'teal': '008080',
         'maroon': '800000', 'lime': '00FF00', 'olive': '808000', 'aqua': '00FFFF', 'fuchsia': 'FF00FF',
         'brown': 'A52A2A', 'pink': 'FFC0CB', 'gold': 'FFD700', 'lightgray': 'D3D3D3', 'lightgrey': 'D3D3D3',
         'darkgray': 'A9A9A9', 'darkgrey': 'A9A9A9', 'transparent': None}


# ---- small helpers ---------------------------------------------------------------------------------------------
def local(tag):
    return etree.QName(tag).localname if isinstance(tag, str) else ''


def num(v, default=0.0):
    """Parse a number with an optional unit (px, pt, mm, %). A percent returns the fraction."""
    if v is None:
        return default
    v = str(v).strip()
    m = re.match(r'^([-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?)\s*(%|px|pt|mm|cm|in|em)?$', v)
    if not m:
        return default
    x = float(m.group(1))
    u = m.group(2)
    if u == '%':
        return x / 100.0
    return x * {'mm': 72 / 25.4, 'cm': 72 / 2.54, 'in': 72.0}.get(u, 1.0)


def mat_mul(m, n):
    """Return m x n: apply n first, then m. A matrix is (a, b, c, d, e, f) as in SVG."""
    return (m[0] * n[0] + m[2] * n[1], m[1] * n[0] + m[3] * n[1],
            m[0] * n[2] + m[2] * n[3], m[1] * n[2] + m[3] * n[3],
            m[0] * n[4] + m[2] * n[5] + m[4], m[1] * n[4] + m[3] * n[5] + m[5])


def mat_apply(m, x, y):
    return (m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5])


def parse_transform(s):
    m = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)
    if not s:
        return m
    for name, args in re.findall(r'(\w+)\s*\(([^)]*)\)', s):
        a = [float(x) for x in re.findall(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?', args)]
        if name == 'translate':
            t = (1, 0, 0, 1, a[0], a[1] if len(a) > 1 else 0.0)
        elif name == 'scale':
            t = (a[0], 0, 0, a[1] if len(a) > 1 else a[0], 0, 0)
        elif name == 'rotate':
            c, sn = math.cos(math.radians(a[0])), math.sin(math.radians(a[0]))
            t = (c, sn, -sn, c, 0, 0)
            if len(a) >= 3:
                t = mat_mul(mat_mul((1, 0, 0, 1, a[1], a[2]), t), (1, 0, 0, 1, -a[1], -a[2]))
        elif name == 'matrix' and len(a) == 6:
            t = tuple(a)
        elif name == 'skewX':
            t = (1, 0, math.tan(math.radians(a[0])), 1, 0, 0)
        elif name == 'skewY':
            t = (1, math.tan(math.radians(a[0])), 0, 1, 0, 0)
        else:
            continue
        m = mat_mul(m, t)
    return m


def decompose(m):
    """Return (sx, sy, rot, shear_ok, flip). rot in radians, clockwise on screen."""
    sx = math.hypot(m[0], m[1])
    sy = math.hypot(m[2], m[3])
    det = m[0] * m[3] - m[1] * m[2]
    rot = math.atan2(m[1], m[0])
    dot = (m[0] * m[2] + m[1] * m[3]) / (sx * sy) if sx and sy else 0.0
    return sx, sy, rot, abs(dot) < 1e-4, det < 0


def parse_color(v):
    """Return (RRGGBB, alpha) or None for none/transparent. alpha in 0..1."""
    if v is None:
        return None
    v = v.strip()
    if not v or v in ('none', 'transparent'):
        return None
    if v.startswith('#'):
        h = v[1:]
        if len(h) in (3, 4):
            h = ''.join(c * 2 for c in h)
        if len(h) == 8:
            return h[:6].upper(), int(h[6:], 16) / 255.0
        if len(h) == 6:
            return h.upper(), 1.0
        return None
    m = re.match(r'rgba?\(([^)]*)\)', v)
    if m:
        parts = [p.strip() for p in m.group(1).replace('/', ',').split(',')]
        rgb = []
        for p in parts[:3]:
            rgb.append(round(float(p[:-1]) * 2.55) if p.endswith('%') else round(float(p)))
        a = float(parts[3]) if len(parts) > 3 else 1.0
        return '%02X%02X%02X' % tuple(max(0, min(255, c)) for c in rgb), a
    c = NAMED.get(v.lower(), 'MISSING')
    if c == 'MISSING':
        return ('000000', 1.0)
    return (c, 1.0) if c else None


def url_id(v):
    m = re.match(r'\s*url\(\s*[\'"]?#([^)\'"]+)[\'"]?\s*\)', v or '')
    return m.group(1) if m else None


# ---- font metrics (width only; the line pitch of PowerPoint does not depend on the font) -----------------------
_FONT_FILES = {
    ('arial', False, False): ['/System/Library/Fonts/Supplemental/Arial.ttf', 'C:/Windows/Fonts/arial.ttf', '/usr/share/fonts/truetype/msttcorefonts/Arial.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'],
    ('arial', True, False): ['/System/Library/Fonts/Supplemental/Arial Bold.ttf', 'C:/Windows/Fonts/arialbd.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'],
    ('arial', False, True): ['/System/Library/Fonts/Supplemental/Arial Italic.ttf', 'C:/Windows/Fonts/ariali.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf'],
    ('arial', True, True): ['/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf', 'C:/Windows/Fonts/arialbi.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf'],
    ('cambria', False, False): ['/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Cambria.ttc', 'C:/Windows/Fonts/cambria.ttc'],
    ('cambria', True, False): ['/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Cambriab.ttf', 'C:/Windows/Fonts/cambriab.ttf'],
    ('cambria', False, True): ['/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Cambriai.ttf', 'C:/Windows/Fonts/cambriai.ttf'],
    ('cambria', True, True): ['/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Cambriaz.ttf', 'C:/Windows/Fonts/cambriaz.ttf'],
    ('times new roman', False, False): ['/System/Library/Fonts/Supplemental/Times New Roman.ttf', 'C:/Windows/Fonts/times.ttf'],
    ('times new roman', True, False): ['/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf'],
    ('times new roman', False, True): ['/System/Library/Fonts/Supplemental/Times New Roman Italic.ttf'],
    ('times new roman', True, True): ['/System/Library/Fonts/Supplemental/Times New Roman Bold Italic.ttf'],
    ('latin modern roman', False, False): [str(HERE.parent / 'fonts' / 'lmroman10-regular.otf')],
    ('latin modern roman', True, False): [str(HERE.parent / 'fonts' / 'lmroman10-bold.otf')],
    ('latin modern roman', False, True): [str(HERE.parent / 'fonts' / 'lmroman10-italic.otf')],
    ('latin modern roman', True, True): [str(HERE.parent / 'fonts' / 'lmroman10-italic.otf')],
}
_font_cache = {}


def _load_font(name, bold, italic):
    key = (name.lower(), bold, italic)
    if key in _font_cache:
        return _font_cache[key]
    f = None
    try:
        from PIL import ImageFont
        for path in _FONT_FILES.get(key, []):
            if os.path.exists(path):
                f = ImageFont.truetype(path, 1000)
                break
    except Exception:
        f = None
    _font_cache[key] = f
    return f


def text_width(s, family, size, bold, italic):
    """Width of s in user units. Falls back to a mean glyph width when the font file is missing."""
    f = _load_font(family, bold, italic)
    if f is not None:
        try:
            return f.getlength(s) * size / 1000.0
        except Exception:
            pass
    serif = family.lower() in ('cambria', 'times new roman', 'latin modern roman')
    return len(s) * size * (0.5 if serif else 0.55) * (1.05 if bold else 1.0)


def first_family(v):
    if not v:
        return 'Arial'
    first = v.split(',')[0].strip().strip('\'"')
    generic = {'sans-serif': 'Arial', 'serif': 'Times New Roman', 'monospace': 'Courier New'}
    return generic.get(first.lower(), first)


# ---- path parsing ----------------------------------------------------------------------------------------------
_NUM = re.compile(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?')


def parse_path(d):
    """Return subpaths as a list of segments in absolute user coordinates.
    A segment is ('M', (x, y)), ('L', (x, y)), ('C', (x1, y1), (x2, y2), (x, y)) or ('Z',)."""
    segs = []
    i, n = 0, len(d)
    cx = cy = sx = sy = 0.0
    last_c = None      # last cubic control point
    last_q = None      # last quadratic control point
    cmd = None
    arity = {'M': 2, 'L': 2, 'H': 1, 'V': 1, 'C': 6, 'S': 4, 'Q': 4, 'T': 2, 'A': 7, 'Z': 0}

    def skip():
        nonlocal i
        while i < n and d[i] in ' \t\r\n,':
            i += 1

    def read_num():
        nonlocal i
        skip()
        m = _NUM.match(d, i)
        if not m:
            raise ValueError('number expected at %d' % i)
        i = m.end()
        return float(m.group(0))

    def read_flag():
        nonlocal i
        skip()
        if i < n and d[i] in '01':
            i += 1
            return int(d[i - 1])
        raise ValueError('flag expected')

    first = True
    while True:
        skip()
        if i >= n:
            break
        ch = d[i]
        if ch.isalpha():
            cmd = ch
            i += 1
            if cmd in 'Zz':
                if segs:
                    segs.append(('Z',))
                cx, cy = sx, sy
                last_c = last_q = None
                continue
        elif cmd is None:
            raise ValueError('path data must start with a command')
        elif cmd in 'Zz':
            raise ValueError('number after Z')
        up = cmd.upper()
        rel = cmd.islower()
        if up == 'M':
            x, y = read_num(), read_num()
            if rel:
                x += cx
                y += cy
            segs.append(('M', (x, y)))
            cx, cy, sx, sy = x, y, x, y
            cmd = 'l' if rel else 'L'
            last_c = last_q = None
        elif up == 'L':
            x, y = read_num(), read_num()
            if rel:
                x += cx
                y += cy
            segs.append(('L', (x, y)))
            cx, cy = x, y
            last_c = last_q = None
        elif up == 'H':
            x = read_num() + (cx if rel else 0)
            segs.append(('L', (x, cy)))
            cx = x
            last_c = last_q = None
        elif up == 'V':
            y = read_num() + (cy if rel else 0)
            segs.append(('L', (cx, y)))
            cy = y
            last_c = last_q = None
        elif up == 'C':
            v = [read_num() for _ in range(6)]
            if rel:
                v = [v[0] + cx, v[1] + cy, v[2] + cx, v[3] + cy, v[4] + cx, v[5] + cy]
            segs.append(('C', (v[0], v[1]), (v[2], v[3]), (v[4], v[5])))
            last_c = (v[2], v[3])
            cx, cy = v[4], v[5]
            last_q = None
        elif up == 'S':
            v = [read_num() for _ in range(4)]
            if rel:
                v = [v[0] + cx, v[1] + cy, v[2] + cx, v[3] + cy]
            c1 = (2 * cx - last_c[0], 2 * cy - last_c[1]) if last_c else (cx, cy)
            segs.append(('C', c1, (v[0], v[1]), (v[2], v[3])))
            last_c = (v[0], v[1])
            cx, cy = v[2], v[3]
            last_q = None
        elif up == 'Q':
            v = [read_num() for _ in range(4)]
            if rel:
                v = [v[0] + cx, v[1] + cy, v[2] + cx, v[3] + cy]
            segs.append(quad_to_cubic((cx, cy), (v[0], v[1]), (v[2], v[3])))
            last_q = (v[0], v[1])
            cx, cy = v[2], v[3]
            last_c = None
        elif up == 'T':
            x, y = read_num(), read_num()
            if rel:
                x += cx
                y += cy
            q = (2 * cx - last_q[0], 2 * cy - last_q[1]) if last_q else (cx, cy)
            segs.append(quad_to_cubic((cx, cy), q, (x, y)))
            last_q = q
            cx, cy = x, y
            last_c = None
        elif up == 'A':
            rx, ry, rot = read_num(), read_num(), read_num()
            fa, fs = read_flag(), read_flag()
            x, y = read_num(), read_num()
            if rel:
                x += cx
                y += cy
            segs.extend(arc_to_cubics((cx, cy), rx, ry, rot, fa, fs, (x, y)))
            cx, cy = x, y
            last_c = last_q = None
        else:
            raise ValueError('unknown command ' + cmd)
        first = False
    return segs


def quad_to_cubic(p0, q, p1):
    return ('C', (p0[0] + 2.0 / 3 * (q[0] - p0[0]), p0[1] + 2.0 / 3 * (q[1] - p0[1])),
            (p1[0] + 2.0 / 3 * (q[0] - p1[0]), p1[1] + 2.0 / 3 * (q[1] - p1[1])), p1)


def arc_to_cubics(p0, rx, ry, phi_deg, fa, fs, p1):
    """Convert an SVG endpoint arc to cubic Bezier segments (at most 90 degrees each)."""
    x1, y1 = p0
    x2, y2 = p1
    if (x1, y1) == (x2, y2):
        return []
    rx, ry = abs(rx), abs(ry)
    if rx == 0 or ry == 0:
        return [('L', p1)]
    phi = math.radians(phi_deg)
    cp, sp = math.cos(phi), math.sin(phi)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    lam = x1p ** 2 / rx ** 2 + y1p ** 2 / ry ** 2
    if lam > 1:
        s = math.sqrt(lam)
        rx, ry = rx * s, ry * s
    num_ = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    co = math.sqrt(max(0.0, num_ / den)) if den else 0.0
    if fa == fs:
        co = -co
    cxp, cyp = co * rx * y1p / ry, -co * ry * x1p / rx
    cx = cp * cxp - sp * cyp + (x1 + x2) / 2
    cy = sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a

    th1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dth = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dth > 0:
        dth -= 2 * math.pi
    elif fs and dth < 0:
        dth += 2 * math.pi
    nseg = max(1, int(math.ceil(abs(dth) / (math.pi / 2) - 1e-9)))
    delta = dth / nseg
    t = 4.0 / 3 * math.tan(delta / 4)
    out = []
    a = th1
    for k in range(nseg):
        b = a + delta
        ca, sa, cb, sb = math.cos(a), math.sin(a), math.cos(b), math.sin(b)

        def pt(ex, ey):
            return (cx + cp * rx * ex - sp * ry * ey, cy + sp * rx * ex + cp * ry * ey)

        c1 = pt(ca - t * sa, sa + t * ca)
        c2 = pt(cb + t * sb, sb - t * cb)
        end = p1 if k == nseg - 1 else pt(cb, sb)
        out.append(('C', c1, c2, end))
        a = b
    return out


def points_to_segs(pts, close):
    nums = [float(v) for v in _NUM.findall(pts or '')]
    xy = list(zip(nums[0::2], nums[1::2]))
    if not xy:
        return []
    segs = [('M', xy[0])] + [('L', p) for p in xy[1:]]
    if close:
        segs.append(('Z',))
    return segs


def segs_bbox(segs):
    """Bounding box (x0, y0, x1, y1) of the drawn curve, by sampling each cubic."""
    xs, ys = [], []
    cur = None
    for s in segs:
        if s[0] in 'ML':
            xs.append(s[1][0])
            ys.append(s[1][1])
            cur = s[1]
        elif s[0] == 'C':
            p0 = cur
            for k in range(1, 9):
                t = k / 8.0
                u = 1 - t
                xs.append(u ** 3 * p0[0] + 3 * u * u * t * s[1][0] + 3 * u * t * t * s[2][0] + t ** 3 * s[3][0])
                ys.append(u ** 3 * p0[1] + 3 * u * u * t * s[1][1] + 3 * u * t * t * s[2][1] + t ** 3 * s[3][1])
            cur = s[3]
    if not xs:
        return None
    return min(xs), min(ys), max(xs), max(ys)


# ---- clipping against a convex polygon (PowerPoint has no clip path, so the geometry itself is cut) -------------
def flatten(segs, steps=12):
    """Subpaths of a segment list as (points, closed) with curves cut into short lines."""
    subs = []
    pts, cur, start = [], None, None
    for s in segs:
        if s[0] == 'M':
            if len(pts) > 1:
                subs.append((pts, False))
            pts, cur, start = [s[1]], s[1], s[1]
        elif s[0] == 'L':
            pts.append(s[1])
            cur = s[1]
        elif s[0] == 'C':
            p0 = cur
            for k in range(1, steps + 1):
                t = k / float(steps)
                u = 1 - t
                pts.append((u ** 3 * p0[0] + 3 * u * u * t * s[1][0] + 3 * u * t * t * s[2][0] + t ** 3 * s[3][0],
                            u ** 3 * p0[1] + 3 * u * u * t * s[1][1] + 3 * u * t * t * s[2][1] + t ** 3 * s[3][1]))
            cur = s[3]
        elif s[0] == 'Z':
            if len(pts) > 1:
                subs.append((pts, True))
            pts, cur = [], start
            if start is not None:
                pts = [start]
    if len(pts) > 1:
        subs.append((pts, False))
    return subs


def poly_area(p):
    return 0.5 * sum(p[i][0] * p[(i + 1) % len(p)][1] - p[(i + 1) % len(p)][0] * p[i][1] for i in range(len(p)))


def is_convex(p):
    n = len(p)
    sign = 0
    for i in range(n):
        a, b, c = p[i], p[(i + 1) % n], p[(i + 2) % n]
        cr = (b[0] - a[0]) * (c[1] - b[1]) - (b[1] - a[1]) * (c[0] - b[0])
        if abs(cr) < 1e-9:
            continue
        s = 1 if cr > 0 else -1
        if sign and s != sign:
            return False
        sign = s
    return True


def inside(p, clip, eps=1e-6):
    """True when the point p is inside the convex polygon clip (counter-clockwise or clockwise)."""
    sgn = 1 if poly_area(clip) > 0 else -1
    n = len(clip)
    for i in range(n):
        a, b = clip[i], clip[(i + 1) % n]
        if sgn * ((b[0] - a[0]) * (p[1] - a[1]) - (b[1] - a[1]) * (p[0] - a[0])) < -eps:
            return False
    return True


def clip_polygon(subject, clip):
    """Sutherland-Hodgman. Both polygons are lists of points, clip is convex."""
    sgn = 1 if poly_area(clip) > 0 else -1
    out = list(subject)
    n = len(clip)
    for i in range(n):
        a, b = clip[i], clip[(i + 1) % n]
        if not out:
            break
        inp, out = out, []

        def side(p):
            return sgn * ((b[0] - a[0]) * (p[1] - a[1]) - (b[1] - a[1]) * (p[0] - a[0]))

        for j in range(len(inp)):
            p, q = inp[j], inp[(j + 1) % len(inp)]
            sp, sq = side(p), side(q)
            if sp >= 0:
                out.append(p)
            if (sp >= 0) != (sq >= 0):
                t = sp / (sp - sq)
                out.append((p[0] + t * (q[0] - p[0]), p[1] + t * (q[1] - p[1])))
    return out


def clip_polyline(pts, closed, clip):
    """Cyrus-Beck per segment. Returns a list of polylines that lie inside the convex polygon clip."""
    sgn = 1 if poly_area(clip) > 0 else -1
    seq = list(pts) + ([pts[0]] if closed else [])
    pieces = []
    cur = []
    n = len(clip)
    for i in range(len(seq) - 1):
        p, q = seq[i], seq[i + 1]
        t0, t1 = 0.0, 1.0
        dx, dy = q[0] - p[0], q[1] - p[1]
        ok = True
        for j in range(n):
            a, b = clip[j], clip[(j + 1) % n]
            ex, ey = b[0] - a[0], b[1] - a[1]
            # inside when sgn * cross(e, x - a) >= 0
            num_ = sgn * (ex * (p[1] - a[1]) - ey * (p[0] - a[0]))
            den = sgn * (ex * dy - ey * dx)
            if abs(den) < 1e-12:
                if num_ < -1e-9:
                    ok = False
                    break
                continue
            t = -num_ / den
            if den > 0:
                t0 = max(t0, t)
            else:
                t1 = min(t1, t)
            if t0 > t1:
                ok = False
                break
        if not ok:
            if len(cur) > 1:
                pieces.append(cur)
            cur = []
            continue
        a_ = (p[0] + t0 * dx, p[1] + t0 * dy)
        b_ = (p[0] + t1 * dx, p[1] + t1 * dy)
        if cur and abs(cur[-1][0] - a_[0]) < 1e-7 and abs(cur[-1][1] - a_[1]) < 1e-7:
            cur.append(b_)
        else:
            if len(cur) > 1:
                pieces.append(cur)
            cur = [a_, b_]
        if t1 < 1.0:
            pieces.append(cur)
            cur = []
    if len(cur) > 1:
        pieces.append(cur)
    return pieces


def rrect_segs(x, y, w, h, rx, ry):
    k = 0.5522847498
    return [('M', (x + rx, y)), ('L', (x + w - rx, y)),
            ('C', (x + w - rx + k * rx, y), (x + w, y + ry - k * ry), (x + w, y + ry)), ('L', (x + w, y + h - ry)),
            ('C', (x + w, y + h - ry + k * ry), (x + w - rx + k * rx, y + h), (x + w - rx, y + h)), ('L', (x + rx, y + h)),
            ('C', (x + rx - k * rx, y + h), (x, y + h - ry + k * ry), (x, y + h - ry)), ('L', (x, y + ry)),
            ('C', (x, y + ry - k * ry), (x + rx - k * rx, y), (x + rx, y)), ('Z',)]


def mat_inv(m):
    det = m[0] * m[3] - m[1] * m[2]
    if abs(det) < 1e-12:
        return None
    a, b, c, d = m[3] / det, -m[1] / det, -m[2] / det, m[0] / det
    return (a, b, c, d, -(a * m[4] + c * m[5]), -(b * m[4] + d * m[5]))


# ---- the converter ---------------------------------------------------------------------------------------------
class Converter:
    def __init__(self, root, slide_part, slide_w, slide_h):
        self.root = root
        self.slide_part = slide_part
        self.defs = {}
        for el in root.iter():
            if isinstance(el.tag, str) and el.get('id'):
                self.defs[el.get('id')] = el
        self.next_id = 2
        self.stats = Counter()
        self.warn = Counter()
        self.approx = Counter()
        vb = root.get('viewBox')
        if vb:
            vx, vy, vw, vh = [float(v) for v in re.split(r'[\s,]+', vb.strip())]
        else:
            vx = vy = 0.0
            vw, vh = num(root.get('width'), 100.0), num(root.get('height'), 100.0)
        self.vb = (vx, vy, vw, vh)
        # Fit the figure inside the slide, centred, never stretched. A 16:9 figure fills the slide.
        self.fit = fit_box(vw, vh, slide_w, slide_h)
        fx, fy, fw, fh = self.fit
        self.k = fw / vw
        self.m0 = (self.k, 0.0, 0.0, fh / vh, fx - vx * self.k, fy - vy * fh / vh)

    # -- ids and names
    def new_id(self):
        self.next_id += 1
        return self.next_id

    def name_of(self, el, fallback):
        for a in ('data-name', 'id', 'data-part', 'data-panel', 'data-role'):
            v = el.get(a)
            if v:
                return v
        return fallback

    # -- style
    def style_of(self, el, parent):
        st = dict(parent)
        for a in INHERITED:
            v = el.get(a)
            if v is not None:
                st[a] = v.strip()
        css = el.get('style')
        if css:
            for part in css.split(';'):
                if ':' in part:
                    k, v = part.split(':', 1)
                    k = k.strip()
                    if k in INHERITED:
                        st[k] = v.strip()
        st['_opacity'] = parent.get('_opacity', 1.0) * num(el.get('opacity'), 1.0)
        for a in ('filter', 'clip-path', 'mask'):
            v = el.get(a)
            if v and v != 'none':
                st['_' + a] = v
        return st

    # -- fills and strokes
    def solid_xml(self, rgb, alpha):
        if alpha < 0.999:
            return '<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>' % (rgb, round(alpha * 100000))
        return '<a:srgbClr val="%s"/>' % rgb

    def stops_of(self, g):
        seen = set()
        while g is not None and id(g) not in seen:
            seen.add(id(g))
            stops = [s for s in g if local(s.tag) == 'stop']
            if stops:
                out = []
                for s in stops:
                    css = {}
                    if s.get('style'):
                        for part in s.get('style').split(';'):
                            if ':' in part:
                                k, v = part.split(':', 1)
                                css[k.strip()] = v.strip()
                    col = parse_color(css.get('stop-color', s.get('stop-color', '#000000'))) or ('000000', 1.0)
                    o = num(css.get('offset', s.get('offset')), 0.0)
                    op = num(css.get('stop-opacity', s.get('stop-opacity')), 1.0)
                    out.append((max(0.0, min(1.0, o)), col[0], col[1] * op))
                return out
            href = g.get('{%s}href' % XLINK) or g.get('href')
            g = self.defs.get(href[1:]) if href and href.startswith('#') else None
        return []

    def color_at(self, stops, t):
        if t <= stops[0][0]:
            return stops[0][1], stops[0][2]
        if t >= stops[-1][0]:
            return stops[-1][1], stops[-1][2]
        for (o0, c0, a0), (o1, c1, a1) in zip(stops, stops[1:]):
            if o0 <= t <= o1:
                f = 0.0 if o1 == o0 else (t - o0) / (o1 - o0)
                rgb = ''.join('%02X' % round(int(c0[i:i + 2], 16) * (1 - f) + int(c1[i:i + 2], 16) * f) for i in (0, 2, 4))
                return rgb, a0 * (1 - f) + a1 * f
        return stops[-1][1], stops[-1][2]

    def grad_xml(self, gid, alpha_mul, geo):
        """Return a <a:gradFill> for the gradient gid, or None. geo describes the shape box in slide EMU."""
        g = self.defs.get(gid)
        if g is None:
            self.warn['gradient: unknown id'] += 1
            return None
        kind = local(g.tag)
        if kind == 'pattern':
            self.approx['pattern fill: solid fallback'] += 1
            for ch in g:
                if local(ch.tag) in ('rect', 'path', 'circle'):
                    c = parse_color(ch.get('fill'))
                    if c:
                        return '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[0], c[1] * alpha_mul)
            return None
        if kind not in ('linearGradient', 'radialGradient'):
            self.warn['fill: unsupported paint server %s' % kind] += 1
            return None
        stops = self.stops_of(g)
        if not stops:
            return None
        if len(stops) == 1:
            return '<a:solidFill>%s</a:solidFill>' % self.solid_xml(stops[0][1], stops[0][2] * alpha_mul)
        if g.get('gradientTransform'):
            self.approx['gradientTransform ignored'] += 1
        units = g.get('gradientUnits', 'objectBoundingBox')
        ux0, uy0, uw, uh = geo['ubox']
        M = geo['M']
        # Local frame of the shape: undo the rotation about the box centre, then measure from the box corner.
        bx, by, bw, bh, rot = geo['x'], geo['y'], geo['w'], geo['h'], geo['rot']
        ccx, ccy = bx + bw / 2.0, by + bh / 2.0
        cr, sr = math.cos(-rot), math.sin(-rot)

        def to_local(ux, uy):
            if units == 'objectBoundingBox':
                ux, uy = ux0 + ux * uw, uy0 + uy * uh
            sx_, sy_ = mat_apply(M, ux, uy)
            dx, dy = sx_ - ccx, sy_ - ccy
            return (cr * dx - sr * dy + bw / 2.0, sr * dx + cr * dy + bh / 2.0)

        def getn(name, default):
            v = g.get(name)
            if v is None:
                return default
            x = num(v, default)
            return x

        gs = []
        if kind == 'linearGradient':
            p1 = to_local(getn('x1', 0.0), getn('y1', 0.0))
            p2 = to_local(getn('x2', 1.0), getn('y2', 0.0))
            dx, dy = p2[0] - p1[0], p2[1] - p1[1]
            dd = dx * dx + dy * dy
            if dd < 1e-9:
                c = stops[-1]
                return '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[1], c[2] * alpha_mul)
            ts = [((cx_ - p1[0]) * dx + (cy_ - p1[1]) * dy) / dd for cx_, cy_ in ((0, 0), (bw, 0), (0, bh), (bw, bh))]
            tmin, tmax = min(ts), max(ts)
            span = tmax - tmin
            if span < 1e-9:
                c = stops[-1]
                return '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[1], c[2] * alpha_mul)
            pts = [(0.0,) + self.color_at(stops, tmin)]
            for o, c, a in stops:
                pos = (o - tmin) / span
                if 0.0 < pos < 1.0:
                    pts.append((pos, c, a))
            pts.append((1.0,) + self.color_at(stops, tmax))
            angle = math.degrees(math.atan2(dy, dx)) % 360.0
            for pos, c, a in pts:
                gs.append('<a:gs pos="%d">%s</a:gs>' % (round(pos * 100000), self.solid_xml(c, a * alpha_mul)))
            return ('<a:gradFill rotWithShape="1"><a:gsLst>%s</a:gsLst><a:lin ang="%d" scaled="0"/></a:gradFill>'
                    % (''.join(gs), round(angle * 60000)))
        # radial
        cxf, cyf = getn('cx', 0.5), getn('cy', 0.5)
        rf = getn('r', 0.5)
        fxf, fyf = getn('fx', cxf), getn('fy', cyf)
        if units != 'objectBoundingBox':
            def frac(x, y):
                return ((x - ux0) / uw if uw else 0.5, (y - uy0) / uh if uh else 0.5)
            cxf, cyf = frac(cxf, cyf)
            fxf, fyf = frac(fxf, fyf)
            rf = rf / uw if uw else 0.5
        # PowerPoint draws the circle from the focus (fillToRect) out to the box. Scale the stops by r.
        scale = rf / RADIAL_SCALE
        pts = []
        for o, c, a in stops:
            pts.append((min(1.0, o * scale), c, a))
        if pts[-1][0] < 1.0:
            pts.append((1.0, pts[-1][1], pts[-1][2]))
        for pos, c, a in pts:
            gs.append('<a:gs pos="%d">%s</a:gs>' % (round(pos * 100000), self.solid_xml(c, a * alpha_mul)))
        l, t = round(max(0.0, min(1.0, fxf)) * 100000), round(max(0.0, min(1.0, fyf)) * 100000)
        return ('<a:gradFill rotWithShape="1"><a:gsLst>%s</a:gsLst><a:path path="circle"><a:fillToRect l="%d" t="%d" r="%d" b="%d"/></a:path></a:gradFill>'
                % (''.join(gs), l, t, 100000 - l, 100000 - t))

    def fill_xml(self, st, geo, closed=True):
        v = st.get('fill', '#000000')
        am = st['_opacity'] * num(st.get('fill-opacity'), 1.0)
        gid = url_id(v)
        if gid:
            x = self.grad_xml(gid, am, geo)
            return x or '<a:noFill/>'
        c = parse_color(v)
        if c is None:
            return '<a:noFill/>'
        return '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[0], c[1] * am)

    def line_xml(self, st, scale, geo, ends=''):
        v = st.get('stroke', 'none')
        sw = num(st.get('stroke-width'), 1.0) * scale
        gid = url_id(v)
        c = None if gid else parse_color(v)
        if (c is None and not gid) or sw <= 0:
            return '<a:ln><a:noFill/></a:ln>'
        am = st['_opacity'] * num(st.get('stroke-opacity'), 1.0)
        if gid:
            fill = self.grad_xml(gid, am, geo) or '<a:noFill/>'
        else:
            fill = '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[0], c[1] * am)
        cap = {'round': 'rnd', 'square': 'sq'}.get(st.get('stroke-linecap', 'butt'), 'flat')
        join = st.get('stroke-linejoin', 'miter')
        jx = {'round': '<a:round/>', 'bevel': '<a:bevel/>'}.get(join, '<a:miter lim="%d"/>' % round(num(st.get('stroke-miterlimit'), 4.0) * 100000))
        dash = ''
        da = st.get('stroke-dasharray')
        if da and da != 'none':
            vals = [float(x) for x in _NUM.findall(da)]
            if vals and sum(vals) > 0:
                if len(vals) % 2:
                    vals = vals * 2
                sw_u = sw / scale if scale else 1.0
                ds = []
                for i in range(0, len(vals), 2):
                    d_ = vals[i] / sw_u * 100000
                    g_ = vals[i + 1] / sw_u * 100000
                    if d_ < 100000:
                        # PowerPoint draws a dash shorter than one line width as a solid line. Use a flat-capped
                        # dash of one line width and keep the period.
                        g_ = max(1000, d_ + g_ - 100000)
                        d_ = 100000
                        cap = 'flat'
                    ds.append('<a:ds d="%d" sp="%d"/>' % (max(1, round(d_)), max(1, round(g_))))
                dash = '<a:custDash>%s</a:custDash>' % ''.join(ds)
        return '<a:ln w="%d" cap="%s">%s%s%s%s</a:ln>' % (max(1, round(sw)), cap, fill, dash, jx, ends)

    def effect_xml(self, st, scale):
        f = st.get('_filter')
        if not f:
            return ''
        fid = url_id(f)
        flt = self.defs.get(fid) if fid else None
        if flt is None:
            self.warn['filter: unknown id'] += 1
            return ''
        kids = [local(c.tag) for c in flt]
        if 'feDropShadow' in kids:
            ds = [c for c in flt if local(c.tag) == 'feDropShadow'][0]
            dx, dy = num(ds.get('dx'), 2.0) * scale, num(ds.get('dy'), 2.0) * scale
            sd = num(ds.get('stdDeviation'), 2.0) * scale
            col = parse_color(ds.get('flood-color', '#000000')) or ('000000', 1.0)
            op = num(ds.get('flood-opacity'), 1.0) * col[1] * st['_opacity']
            self.stats['drop shadows -> outer shadow'] += 1
            return ('<a:effectLst><a:outerShdw blurRad="%d" dist="%d" dir="%d" algn="ctr" rotWithShape="0">%s</a:outerShdw></a:effectLst>'
                    % (round(sd * 2), round(math.hypot(dx, dy)), round((math.degrees(math.atan2(dy, dx)) % 360) * 60000),
                       self.solid_xml(col[0], op)))
        if kids == ['feGaussianBlur']:
            sd = num(flt[0].get('stdDeviation'), 1.0) * scale
            self.approx['blur filter -> soft edge'] += 1
            return '<a:effectLst><a:softEdge rad="%d"/></a:effectLst>' % max(1, round(sd * 1.5))
        if 'feMerge' in kids:
            self.warn['glow filter (blur + merge) dropped'] += 1
        else:
            self.warn['filter dropped (%s)' % '+'.join(kids)] += 1
        return ''

    # -- shape emitters. Each returns (xml, bbox) with bbox in slide EMU, or None.
    def shape_header(self, el, name, tag='sp', txbox=False):
        sid = self.new_id()
        if tag == 'cxnSp':
            return sid, '<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="%d" name=%s/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>' % (sid, quoteattr(name))
        return sid, '<p:sp><p:nvSpPr><p:cNvPr id="%d" name=%s/><p:cNvSpPr%s/><p:nvPr/></p:nvSpPr>' % (sid, quoteattr(name), ' txBox="1"' if txbox else '')

    def emit_box_shape(self, el, st, M, ubox, prst, adj=None, name=''):
        """Rect, round rect or ellipse as a preset shape. ubox = (x, y, w, h) in user units."""
        x, y, w, h = ubox
        sx, sy, rot, shear_ok, flip = decompose(M)
        if not shear_ok or flip:
            return None
        cx_, cy_ = mat_apply(M, x + w / 2.0, y + h / 2.0)
        W, H = w * sx, h * sy
        geo = dict(x=cx_ - W / 2, y=cy_ - H / 2, w=W, h=H, rot=rot, ubox=ubox, M=M)
        scale = math.sqrt(abs(M[0] * M[3] - M[1] * M[2]))
        sid, head = self.shape_header(el, name or self.name_of(el, prst))
        rot_attr = ' rot="%d"' % round((math.degrees(rot) % 360) * 60000) if abs(rot) > 1e-9 else ''
        av = '<a:avLst/>' if adj is None else '<a:avLst><a:gd name="adj" fmla="val %d"/></a:avLst>' % adj
        xml = (head + '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="%s">%s</a:prstGeom>%s%s%s</p:spPr></p:sp>'
               % (rot_attr, round(geo['x']), round(geo['y']), max(1, round(W)), max(1, round(H)), prst, av,
                  self.fill_xml(st, geo), self.line_xml(st, scale, geo), self.effect_xml(st, scale)))
        self.stats['shapes'] += 1
        return xml, rot_bbox(geo['x'], geo['y'], W, H, rot)

    def emit_line(self, el, st, M):
        x1, y1, x2, y2 = (num(el.get(a)) for a in ('x1', 'y1', 'x2', 'y2'))
        return self.emit_segments(el, st, M, [('M', (x1, y1)), ('L', (x2, y2))], force_line=True)

    def emit_segments(self, el, st, M, segs, force_line=False, ubox_ref=None):
        """Lines as connectors. Anything else as a freeform (custGeom)."""
        if not segs:
            return None
        scale = math.sqrt(abs(M[0] * M[3] - M[1] * M[2]))
        # transform segments to slide EMU
        tsegs = []
        for s in segs:
            if s[0] == 'Z':
                tsegs.append(s)
            else:
                tsegs.append((s[0],) + tuple(mat_apply(M, p[0], p[1]) for p in s[1:]))
        bb = segs_bbox(tsegs)
        if bb is None:
            return None
        x0, y0, x1, y1 = bb
        ub = segs_bbox(segs)
        ubox = ubox_ref or (ub[0], ub[1], ub[2] - ub[0], ub[3] - ub[1])
        geo = dict(x=x0, y=y0, w=x1 - x0, h=y1 - y0, rot=0.0, ubox=ubox, M=M)
        name = self.name_of(el, local(el.tag))
        ends = self.marker_ends(el, st)
        eff = self.effect_xml(st, scale)
        simple_line = (len(tsegs) == 2 and tsegs[0][0] == 'M' and tsegs[1][0] == 'L')
        if simple_line:
            (sx_, sy_), (ex_, ey_) = tsegs[0][1], tsegs[1][1]
            sid, head = self.shape_header(el, name, 'cxnSp')
            flip = (' flipH="1"' if ex_ < sx_ else '') + (' flipV="1"' if ey_ < sy_ else '')
            xml = (head + '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom>%s%s</p:spPr></p:cxnSp>'
                   % (flip, round(x0), round(y0), round(x1 - x0), round(y1 - y0),
                      self.line_xml(st, scale, geo, ends), eff))
            self.stats['shapes'] += 1
            return xml, (x0, y0, x1, y1)
        W, H = max(1.0, x1 - x0), max(1.0, y1 - y0)
        parts = []
        closed_all = True
        for s in tsegs:
            if s[0] == 'M':
                parts.append('<a:moveTo><a:pt x="%d" y="%d"/></a:moveTo>' % (round(s[1][0] - x0), round(s[1][1] - y0)))
            elif s[0] == 'L':
                parts.append('<a:lnTo><a:pt x="%d" y="%d"/></a:lnTo>' % (round(s[1][0] - x0), round(s[1][1] - y0)))
            elif s[0] == 'C':
                parts.append('<a:cubicBezTo>%s</a:cubicBezTo>' % ''.join('<a:pt x="%d" y="%d"/>' % (round(p[0] - x0), round(p[1] - y0)) for p in s[1:]))
            else:
                parts.append('<a:close/>')
        closed = tsegs[-1][0] == 'Z'
        # An open path with a fill is closed by the fill itself in SVG and PowerPoint alike.
        fill = self.fill_xml(st, geo)
        sid, head = self.shape_header(el, name)
        xml = (head + '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
               '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
               '<a:pathLst><a:path w="%d" h="%d">%s</a:path></a:pathLst></a:custGeom>%s%s%s</p:spPr></p:sp>'
               % (round(x0), round(y0), round(W), round(H), round(W), round(H), ''.join(parts), fill,
                  self.line_xml(st, scale, geo, ends), eff))
        if st.get('fill-rule') == 'evenodd':
            self.approx['fill-rule evenodd: PowerPoint decides holes'] += 1
        self.stats['shapes'] += 1
        return xml, (x0, y0, x1, y1)

    def marker_ends(self, el, st):
        out = ''
        for key, tag in (('marker-start', 'headEnd'), ('marker-end', 'tailEnd')):
            v = st.get(key)
            if v and v != 'none':
                self.approx['marker -> triangle line end'] += 1
                out += '<a:%s type="triangle" w="med" len="med"/>' % tag
        return out

    def emit_image(self, el, st, M):
        href = el.get('{%s}href' % XLINK) or el.get('href') or ''
        m = re.match(r'data:image/(png|jpeg|jpg|gif);base64,(.*)', href, re.S)
        if not m:
            self.warn['image: not a raster data URI, dropped'] += 1
            return None
        x, y, w, h = (num(el.get(a)) for a in ('x', 'y', 'width', 'height'))
        sx, sy, rot, shear_ok, flip = decompose(M)
        if not shear_ok or flip:
            self.warn['image: sheared or mirrored, dropped'] += 1
            return None
        data = base64.b64decode(m.group(2))
        _, rid = self.slide_part.get_or_add_image_part(io.BytesIO(data))
        cx_, cy_ = mat_apply(M, x + w / 2.0, y + h / 2.0)
        W, H = w * sx, h * sy
        sid = self.new_id()
        rot_attr = ' rot="%d"' % round((math.degrees(rot) % 360) * 60000) if abs(rot) > 1e-9 else ''
        xml = ('<p:pic><p:nvPicPr><p:cNvPr id="%d" name=%s/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
               '<p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
               '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
               % (sid, quoteattr(self.name_of(el, 'image')), rid, rot_attr, round(cx_ - W / 2), round(cy_ - H / 2), round(W), round(H)))
        self.stats['pictures'] += 1
        return xml, rot_bbox(cx_ - W / 2, cy_ - H / 2, W, H, rot)

    # -- text
    def collect_lines(self, el, st):
        """Split a <text> into lines of runs. Each run keeps its own size, weight, style, colour and baseline offset."""
        size0 = num(st.get('font-size'), 12.0)
        preserve = (el.get('{http://www.w3.org/XML/1998/namespace}space') == 'preserve'
                    or 'white-space: pre' in (el.get('style') or ''))
        x0, y0 = num(el.get('x')), num(el.get('y'))
        lines = []
        state = {'y': 0.0, 'line': None, 'x': x0}

        def newline(x):
            ln = {'x': x, 'runs': [], 'base': None}
            lines.append(ln)
            state['line'] = ln
            return ln

        base_rs = {'size': size0, 'bold': st.get('font-weight', 'normal') in ('bold', '700', '800', '900', 'bolder'),
                   'italic': st.get('font-style') == 'italic', 'fill': st.get('fill', '#000000'),
                   'underline': 'underline' in st.get('text-decoration', ''), 'spc': num(st.get('letter-spacing'), 0.0)}

        def add(s, rs):
            ln = state['line'] or newline(state['x'])
            s = s.replace('\u200b', '')
            if not preserve:
                s = re.sub(r'\s+', ' ', s)
                if not ln['runs'] or all(r['text'] == '' for r in ln['runs']):
                    s = s.lstrip()
            small = rs['size'] < size0 * 0.95
            if ln['base'] is None and not small:
                ln['base'] = state['y']
            ln['runs'].append(dict(rs, text=s, y=state['y']))

        def walk(node, rs, first):
            for att in ('x', 'y'):
                pass
            if node is not el:
                if node.get('x') is not None or (node.get('y') is not None):
                    cur = state['line']
                    if cur is not None and any(r['text'] for r in cur['runs']):
                        newline(num(node.get('x'), state['x']))
                    elif cur is None:
                        newline(num(node.get('x'), state['x']))
                    else:
                        cur['x'] = num(node.get('x'), cur['x'])
                    if node.get('y') is not None:
                        state['y'] = num(node.get('y')) - y0
                if node.get('dy') is not None:
                    state['y'] += num(node.get('dy'))
                if node.get('dx') is not None and abs(num(node.get('dx'))) > 1e-9:
                    self.approx['tspan dx ignored'] += 1
                rs = dict(rs)
                if node.get('font-size'):
                    rs['size'] = num(node.get('font-size'), rs['size'])
                if node.get('font-weight'):
                    rs['bold'] = node.get('font-weight') in ('bold', '700', '800', '900', 'bolder')
                if node.get('font-style'):
                    rs['italic'] = node.get('font-style') == 'italic'
                if node.get('fill'):
                    rs['fill'] = node.get('fill')
            if node.text:
                add(node.text, rs)
            for ch in node:
                if local(ch.tag) == 'tspan':
                    walk(ch, rs, False)
                elif local(ch.tag) == 'textPath':
                    walk(ch, rs, False)
                if ch.tail:
                    add(ch.tail, rs)
            if node is not el and not (node.text or len(node)):
                add('', rs)

        newline(x0)
        walk(el, base_rs, True)
        # drop an empty first line and lines with no content at all
        lines = [ln for ln in lines if any(r['text'] for r in ln['runs'])]
        for ln in lines:
            if ln['base'] is None:
                ln['base'] = ln['runs'][0]['y']
            # strip trailing space of the line
            if not preserve and ln['runs']:
                ln['runs'][-1]['text'] = ln['runs'][-1]['text'].rstrip()
        return lines, size0

    def emit_text(self, el, st, M):
        tp = None
        for ch in el:
            if local(ch.tag) == 'textPath':
                tp = ch
        if tp is not None:
            M, st = self.textpath_frame(el, tp, st, M)
        lines, size0 = self.collect_lines(el, st)
        if not lines:
            return None
        if st.get('_clips'):
            inv = mat_inv(M)
            ax_, ay_ = mat_apply(M, lines[0]['x'], num(el.get('y')))
            if not all(inside((ax_, ay_), poly) for poly in st['_clips']):
                self.warn['text outside its clip path dropped'] += 1
                return None
        if st.get('paint-order', '').startswith('stroke') and st.get('stroke', 'none') != 'none':
            self.approx['text halo (paint-order stroke) dropped'] += 1
        sx, sy, rot, shear_ok, flip = decompose(M)
        if not shear_ok or flip:
            self.approx['sheared or mirrored text drawn unsheared'] += 1
        s = math.sqrt(abs(M[0] * M[3] - M[1] * M[2]))        # EMU per user unit
        family = first_family(st.get('font-family'))
        anchor = st.get('text-anchor', 'start')
        x0, y0 = num(el.get('x')), num(el.get('y'))
        if tp is not None:
            x0, y0 = 0.0, 0.0
        ax = lines[0]['x'] if tp is None else 0.0
        if any(abs(ln['x'] - ax) > 1e-6 for ln in lines):
            self.approx['multi-line text with different x: first x used'] += 1
        base0 = y0 + lines[0]['base']
        # line pitch from the baselines of the lines
        if len(lines) > 1:
            pitch = (lines[-1]['base'] - lines[0]['base']) / (len(lines) - 1)
        else:
            pitch = SINGLE_LINE * size0
        lsize = [max(r['size'] for r in ln['runs'] if r['text']) for ln in lines]
        p = pitch / (SINGLE_LINE * lsize[0])
        asc = (BASE_ASC + ASC_PER_P * (p - 1.0)) * lsize[0]
        top = base0 - asc + TEXT_NUDGE_EMU / s
        widths = []
        for ln in lines:
            w = sum(text_width(r['text'], family, r['size'], r['bold'], r['italic']) for r in ln['runs'])
            widths.append(w)
        Wb = max(widths) + 0.4 * size0
        Hb = pitch * (len(lines) - 1) + SINGLE_LINE * lsize[-1] if len(lines) > 1 else SINGLE_LINE * lsize[0]
        if anchor == 'middle':
            cxl = ax
        elif anchor == 'end':
            cxl = ax - Wb / 2.0
        else:
            cxl = ax + Wb / 2.0
        cyl = top + Hb / 2.0
        csx, csy = mat_apply(M, cxl, cyl)
        W, H = Wb * s, Hb * s
        algn = {'middle': 'ctr', 'end': 'r'}.get(anchor, 'l')
        paras = []
        for li, ln in enumerate(lines):
            runs = []
            for r in ln['runs']:
                if not r['text']:
                    continue
                shift = r['y'] - ln['base']
                size_r = r['size']
                if abs(shift) > 1e-6:
                    size_r = size_r / SCRIPT_AUTO    # PowerPoint shrinks a raised or lowered run by itself
                sz = max(100, round(size_r * s / EMU_PT * 100))
                attrs = ' lang="en-US" sz="%d"' % sz
                if r['bold']:
                    attrs += ' b="1"'
                if r['italic']:
                    attrs += ' i="1"'
                if r['underline']:
                    attrs += ' u="sng"'
                if abs(shift) > 1e-6:
                    attrs += ' baseline="%d"' % round(-shift / ln_size(ln, r) * 100000)
                    self.stats['sub/superscript runs'] += 1
                if r['spc']:
                    attrs += ' spc="%d"' % round(r['spc'] * s / EMU_PT * 100)
                c = parse_color(r['fill']) or ('000000', 1.0)
                fill = '<a:solidFill>%s</a:solidFill>' % self.solid_xml(c[0], c[1] * st['_opacity'] * num(st.get('fill-opacity'), 1.0))
                tf = '<a:latin typeface=%s/><a:cs typeface=%s/>' % (quoteattr(family), quoteattr(family))
                runs.append('<a:r><a:rPr%s dirty="0">%s%s</a:rPr><a:t>%s</a:t></a:r>' % (attrs, fill, tf, escape(r['text'])))
            if li == 0:
                pct = round(p * 100000) if len(lines) > 1 else 100000
            lnspc = '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef>' % (round(p * 100000) if len(lines) > 1 else 100000)
            end_sz = max(100, round(lsize[li] * s / EMU_PT * 100))
            paras.append('<a:p><a:pPr algn="%s">%s</a:pPr>%s<a:endParaRPr lang="en-US" sz="%d" dirty="0"/></a:p>' % (algn, lnspc, ''.join(runs), end_sz))
        content = ' / '.join(''.join(r['text'] for r in ln['runs']) for ln in lines)
        sid = self.new_id()
        name = self.name_of(el, 'Text') if (el.get('data-name') or el.get('id')) else 'Text: ' + content[:40]
        rot_attr = ' rot="%d"' % round((math.degrees(rot) % 360) * 60000) if abs(rot) > 1e-9 else ''
        xml = ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name=%s/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
               '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
               '<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="t"><a:noAutofit/></a:bodyPr><a:lstStyle/>%s</p:txBody></p:sp>'
               % (sid, quoteattr(name), rot_attr, round(csx - W / 2), round(csy - H / 2), max(1, round(W)), max(1, round(H)), ''.join(paras)))
        self.stats['text boxes'] += 1
        return xml, rot_bbox(csx - W / 2, csy - H / 2, W, H, rot)

    def textpath_frame(self, el, tp, st, M):
        """Straight text along the chord of the path (PowerPoint has no text on a path)."""
        self.approx['textPath -> straight text along the chord'] += 1
        href = tp.get('{%s}href' % XLINK) or tp.get('href') or ''
        pth = self.defs.get(href[1:]) if href.startswith('#') else None
        st = dict(st)
        if tp.get('text-anchor'):
            st['text-anchor'] = tp.get('text-anchor')
        if pth is None:
            return M, st
        segs = parse_path(pth.get('d', ''))
        pts = [s[-1] for s in segs if s[0] in 'MLC']
        if len(pts) < 2:
            return M, st
        (ax_, ay_), (bx_, by_) = pts[0], pts[-1]
        off = tp.get('startOffset', '0')
        f = num(off, 0.0) if off.strip().endswith('%') else 0.0
        px, py = ax_ + (bx_ - ax_) * f, ay_ + (by_ - ay_) * f
        ang = math.atan2(by_ - ay_, bx_ - ax_)
        c, sn = math.cos(ang), math.sin(ang)
        # local frame: origin at (px, py), x axis along the chord
        M2 = mat_mul(M, (c, sn, -sn, c, px, py))
        return M2, st

    # -- tree walk
    def walk(self, el, M, parent_st, depth):
        """Return a list of (xml, bbox) for the children of a container."""
        out = []
        for ch in el:
            if not isinstance(ch.tag, str):
                continue
            tag = local(ch.tag)
            if tag in SKIP_TAGS:
                continue
            if ch.get('display') == 'none' or ch.get('visibility') == 'hidden':
                continue
            st = self.style_of(ch, parent_st)
            Mc = mat_mul(M, parse_transform(ch.get('transform')))
            cpid = url_id(ch.get('clip-path'))
            if cpid:
                poly = self.clip_polygon_of(cpid, Mc)
                st = dict(st)
                if poly:
                    st['_clips'] = list(st.get('_clips', [])) + [poly]
                else:
                    self.warn['clip-path dropped (not a single convex shape)'] += 1
            try:
                if tag == 'g' or tag == 'a':
                    out.extend(self.emit_group(ch, st, Mc, depth))
                elif tag == 'use':
                    out.extend(self.emit_use(ch, st, Mc, depth))
                else:
                    r = self.emit_leaf(ch, tag, st, Mc)
                    if r:
                        out.append(r)
            except Exception as e:      # never fail on one bad element
                self.warn['element failed: %s (%s)' % (tag, type(e).__name__)] += 1
                if os.environ.get('SVG2PPTX_DEBUG'):
                    raise
        return out

    def shape_segs(self, el, tag):
        """User-space segments of a basic shape or path, with rounded corners kept."""
        if tag == 'rect':
            x, y, w, h = (num(el.get(a)) for a in ('x', 'y', 'width', 'height'))
            rx, ry = el.get('rx'), el.get('ry')
            rxv = num(rx if rx is not None else ry)
            ryv = num(ry if ry is not None else rx)
            rxv, ryv = min(rxv, w / 2), min(ryv, h / 2)
            return rrect_segs(x, y, w, h, rxv, ryv) if rxv > 0 and ryv > 0 else rect_segs(x, y, w, h)
        if tag == 'circle':
            r = num(el.get('r'))
            return ellipse_segs(num(el.get('cx')), num(el.get('cy')), r, r)
        if tag == 'ellipse':
            return ellipse_segs(num(el.get('cx')), num(el.get('cy')), num(el.get('rx')), num(el.get('ry')))
        if tag == 'line':
            return [('M', (num(el.get('x1')), num(el.get('y1')))), ('L', (num(el.get('x2')), num(el.get('y2'))))]
        if tag in ('polyline', 'polygon'):
            return points_to_segs(el.get('points'), tag == 'polygon')
        if tag == 'path':
            return parse_path(el.get('d', ''))
        return []

    def clip_polygon_of(self, cpid, M):
        """The clip path as one convex polygon in slide EMU, or None when it is not that."""
        cp = self.defs.get(cpid)
        if cp is None or local(cp.tag) != 'clipPath':
            return None
        kids = [c for c in cp if isinstance(c.tag, str) and local(c.tag) in ('rect', 'circle', 'ellipse', 'path', 'polygon')]
        if len(kids) != 1:
            return None
        k = kids[0]
        Mk = mat_mul(mat_mul(M, parse_transform(cp.get('transform'))), parse_transform(k.get('transform')))
        try:
            subs = flatten(self.shape_segs(k, local(k.tag)), 16)
        except ValueError:
            return None
        if len(subs) != 1:
            return None
        poly = [mat_apply(Mk, x, y) for x, y in subs[0][0]]
        if len(poly) > 1 and poly[0] == poly[-1]:
            poly = poly[:-1]
        if len(poly) < 3 or not is_convex(poly):
            return None
        return poly

    def emit_clipped(self, el, tag, st, M):
        """Cut the geometry of a leaf with the convex clip polygons. Returns False when nothing needs cutting."""
        inv = mat_inv(M)
        if inv is None:
            return False
        clips = [[mat_apply(inv, x, y) for x, y in poly] for poly in st['_clips']]
        try:
            segs = self.shape_segs(el, tag)
        except ValueError:
            return False
        subs = flatten(segs, 12)
        if not subs:
            return None
        if all(inside(p, c) for c in clips for pts, _ in subs for p in pts):
            return False
        has_fill = parse_color(st.get('fill', '#000000')) is not None or url_id(st.get('fill', ''))
        has_stroke = (parse_color(st.get('stroke', 'none')) is not None or url_id(st.get('stroke', ''))) and num(st.get('stroke-width'), 1.0) > 0
        closed_shape = tag in ('rect', 'circle', 'ellipse', 'polygon') or any(cl for _, cl in subs)
        ub = segs_bbox(segs)
        ubox = (ub[0], ub[1], ub[2] - ub[0], ub[3] - ub[1])
        items = []
        if has_fill and closed_shape:
            fsegs = []
            for pts, closed in subs:
                poly = pts[:-1] if len(pts) > 1 and pts[0] == pts[-1] else pts
                for c in clips:
                    poly = clip_polygon(poly, c) if poly else []
                if len(poly) >= 3:
                    fsegs += [('M', poly[0])] + [('L', q) for q in poly[1:]] + [('Z',)]
            if fsegs:
                s2 = dict(st)
                s2['stroke'] = 'none'
                r = self.emit_segments(el, s2, M, fsegs, ubox_ref=ubox)
                if r:
                    items.append(r)
        if has_stroke:
            ssegs = []
            for pts, closed in subs:
                pieces = [pts]
                closed_flags = [closed]
                for c in clips:
                    newp = []
                    for pc, cl in zip(pieces, closed_flags):
                        newp += [(q, False) for q in clip_polyline(pc, cl, c)]
                    pieces = [q for q, _ in newp]
                    closed_flags = [False] * len(pieces)
                for pc in pieces:
                    ssegs += [('M', pc[0])] + [('L', q) for q in pc[1:]]
            if ssegs:
                s2 = dict(st)
                s2['fill'] = 'none'
                r = self.emit_segments(el, s2, M, ssegs, ubox_ref=ubox)
                if r:
                    items.append(r)
        self.stats['shapes cut by a clip path'] += 1
        if not items:
            return None
        if len(items) == 1:
            return items[0]
        return self.wrap_group(items, self.name_of(el, tag) + ' (clipped)')

    def wrap_group(self, kids, name):
        x0 = min(k[1][0] for k in kids)
        y0 = min(k[1][1] for k in kids)
        x1 = max(k[1][2] for k in kids)
        y1 = max(k[1][3] for k in kids)
        gid = self.new_id()
        xml = ('<p:grpSp><p:nvGrpSpPr><p:cNvPr id="%d" name=%s/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
               '<p:grpSpPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/><a:chOff x="%d" y="%d"/><a:chExt cx="%d" cy="%d"/></a:xfrm></p:grpSpPr>%s</p:grpSp>'
               % (gid, quoteattr(name), round(x0), round(y0), max(1, round(x1 - x0)), max(1, round(y1 - y0)),
                  round(x0), round(y0), max(1, round(x1 - x0)), max(1, round(y1 - y0)), ''.join(k[0] for k in kids)))
        self.stats['groups'] += 1
        return (xml, (x0, y0, x1, y1))

    def emit_use(self, el, st, M, depth):
        href = el.get('{%s}href' % XLINK) or el.get('href') or ''
        tgt = self.defs.get(href[1:]) if href.startswith('#') else None
        if tgt is None:
            self.warn['use: unknown target'] += 1
            return []
        M = mat_mul(M, (1, 0, 0, 1, num(el.get('x')), num(el.get('y'))))
        holder = etree.Element('{%s}g' % SVGNS)
        # deepcopy, not tostring + fromstring: tostring of an inner element can also write the elements after it
        copy_of_target = copy.deepcopy(tgt)
        copy_of_target.tail = None
        holder.append(copy_of_target)
        return self.walk(holder, M, st, depth)

    def emit_group(self, el, st, M, depth):
        if el.get('clip-path') and url_id(el.get('clip-path')):
            self.stats['clip paths applied by cutting the geometry'] += 1
        if el.get('mask'):
            self.warn['mask dropped'] += 1
        if el.get('opacity') and num(el.get('opacity'), 1.0) < 0.999:
            self.approx['group opacity pushed down to children'] += 1
        if el.get('filter') and el.get('filter') != 'none':
            self.approx['group filter applied to each child'] += 1
        kids = self.walk(el, M, st, depth + 1)
        if not kids:
            return []
        if len(kids) == 1 or depth >= MAX_GROUP_DEPTH:
            if depth >= MAX_GROUP_DEPTH:
                self.stats['groups flattened (deeper than %d)' % MAX_GROUP_DEPTH] += 1
            return kids
        x0 = min(k[1][0] for k in kids)
        y0 = min(k[1][1] for k in kids)
        x1 = max(k[1][2] for k in kids)
        y1 = max(k[1][3] for k in kids)
        gid = self.new_id()
        name = self.name_of(el, 'Group')
        xml = ('<p:grpSp><p:nvGrpSpPr><p:cNvPr id="%d" name=%s/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
               '<p:grpSpPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/><a:chOff x="%d" y="%d"/><a:chExt cx="%d" cy="%d"/></a:xfrm></p:grpSpPr>%s</p:grpSp>'
               % (gid, quoteattr(name), round(x0), round(y0), max(1, round(x1 - x0)), max(1, round(y1 - y0)),
                  round(x0), round(y0), max(1, round(x1 - x0)), max(1, round(y1 - y0)), ''.join(k[0] for k in kids)))
        self.stats['groups'] += 1
        return [(xml, (x0, y0, x1, y1))]

    def emit_leaf(self, el, tag, st, M):
        if st.get('_mask'):
            self.warn['mask dropped'] += 1
        if el.get('clip-path') and url_id(el.get('clip-path')):
            self.stats['clip paths applied by cutting the geometry'] += 1
        if st.get('_clips') and tag in ('rect', 'circle', 'ellipse', 'line', 'polyline', 'polygon', 'path'):
            r = self.emit_clipped(el, tag, st, M)
            if r is not False:
                return r
        if tag == 'rect':
            x, y, w, h = (num(el.get(a)) for a in ('x', 'y', 'width', 'height'))
            if w <= 0 or h <= 0:
                return None
            rx, ry = el.get('rx'), el.get('ry')
            if rx is None and ry is None:
                return self.emit_box_shape(el, st, M, (x, y, w, h), 'rect') or self.emit_segments(el, st, M, rect_segs(x, y, w, h))
            rxv = num(rx if rx is not None else ry)
            ryv = num(ry if ry is not None else rx)
            rxv, ryv = min(rxv, w / 2), min(ryv, h / 2)
            if rxv <= 0 or ryv <= 0:
                return self.emit_box_shape(el, st, M, (x, y, w, h), 'rect') or self.emit_segments(el, st, M, rect_segs(x, y, w, h))
            if abs(rxv - ryv) > 0.02 * max(rxv, ryv):
                self.approx['rounded rect with elliptic corners: circular corners used'] += 1
            adj = round(min(50000, min(rxv, ryv) / min(w, h) * 100000))
            return self.emit_box_shape(el, st, M, (x, y, w, h), 'roundRect', adj) or self.emit_segments(el, st, M, rect_segs(x, y, w, h))
        if tag == 'circle':
            cx_, cy_, r = num(el.get('cx')), num(el.get('cy')), num(el.get('r'))
            if r <= 0:
                return None
            return (self.emit_box_shape(el, st, M, (cx_ - r, cy_ - r, 2 * r, 2 * r), 'ellipse')
                    or self.emit_segments(el, st, M, ellipse_segs(cx_, cy_, r, r)))
        if tag == 'ellipse':
            cx_, cy_, rx, ry = (num(el.get(a)) for a in ('cx', 'cy', 'rx', 'ry'))
            if rx <= 0 or ry <= 0:
                return None
            return (self.emit_box_shape(el, st, M, (cx_ - rx, cy_ - ry, 2 * rx, 2 * ry), 'ellipse')
                    or self.emit_segments(el, st, M, ellipse_segs(cx_, cy_, rx, ry)))
        if tag == 'line':
            return self.emit_line(el, st, M)
        if tag in ('polyline', 'polygon'):
            segs = points_to_segs(el.get('points'), tag == 'polygon')
            if len(segs) < 2:
                return None
            return self.emit_segments(el, st, M, segs)
        if tag == 'path':
            try:
                segs = parse_path(el.get('d', ''))
            except ValueError:
                self.warn['path: bad data dropped'] += 1
                return None
            if not segs:
                return None
            return self.emit_segments(el, st, M, segs)
        if tag == 'text':
            return self.emit_text(el, st, M)
        if tag == 'image':
            return self.emit_image(el, st, M)
        self.warn['unknown element <%s> skipped' % tag] += 1
        return None

    def convert(self):
        st0 = {'fill': '#000000', '_opacity': 1.0, 'font-family': 'Arial', 'font-size': '12'}
        return self.walk(self.root, self.m0, st0, 0)


def ln_size(ln, run):
    """Font size that a PowerPoint baseline offset is a percentage of. Test: the size of the line's base text."""
    return max(r['size'] for r in ln['runs'] if r['text'])


def rect_segs(x, y, w, h):
    return [('M', (x, y)), ('L', (x + w, y)), ('L', (x + w, y + h)), ('L', (x, y + h)), ('Z',)]


def ellipse_segs(cx, cy, rx, ry):
    k = 0.5522847498
    return [('M', (cx + rx, cy)),
            ('C', (cx + rx, cy + k * ry), (cx + k * rx, cy + ry), (cx, cy + ry)),
            ('C', (cx - k * rx, cy + ry), (cx - rx, cy + k * ry), (cx - rx, cy)),
            ('C', (cx - rx, cy - k * ry), (cx - k * rx, cy - ry), (cx, cy - ry)),
            ('C', (cx + k * rx, cy - ry), (cx + rx, cy - k * ry), (cx + rx, cy)), ('Z',)]


def rot_bbox(x, y, w, h, rot):
    if abs(rot) < 1e-9:
        return (x, y, x + w, y + h)
    cx, cy = x + w / 2.0, y + h / 2.0
    c, s = math.cos(rot), math.sin(rot)
    xs, ys = [], []
    for px, py in ((-w / 2, -h / 2), (w / 2, -h / 2), (w / 2, h / 2), (-w / 2, h / 2)):
        xs.append(cx + c * px - s * py)
        ys.append(cy + s * px + c * py)
    return (min(xs), min(ys), max(xs), max(ys))


def fit_box(fig_w, fig_h, slide_w, slide_h):
    """Position (x, y, w, h in EMU) of a figure of shape fig_w : fig_h on the slide. The same rule as svg2pptx.py:
    a 16:9 figure fills the slide, any other figure is fitted inside it and centred, never stretched."""
    k = min(slide_w / fig_w, slide_h / fig_h)
    if abs(fig_w / fig_h - 16 / 9) < 0.01:
        w, h = slide_w, slide_h
    else:
        w, h = round(fig_w * k), round(fig_h * k)
    return ((slide_w - w) // 2, (slide_h - h) // 2, w, h)


def build_slide(prs, svg_path, report=True):
    """Add one slide with the native objects of svg_path. Returns (stats, warnings, approximations)."""
    from pptx.oxml import parse_xml
    root = etree.parse(str(svg_path)).getroot()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    conv = Converter(root, slide.part, prs.slide_width, prs.slide_height)
    items = conv.convert()
    tree = slide.shapes._spTree
    for xml, _ in items:
        tree.append(parse_xml('<p:wrap %s>%s</p:wrap>' % (NS_DECL, xml))[0])
    # white page background
    cSld = slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld')
    bg = parse_xml('<p:bg %s><p:bgPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>' % NS_DECL)
    cSld.insert(0, bg)
    return conv.stats, conv.warn, conv.approx


def convert(out, svgs, quiet=False):
    from pptx import Presentation
    from pptx.util import Mm
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(SLIDE_W_MM), Mm(SLIDE_H_MM)
    allres = []
    for i, svg in enumerate(svgs, 1):
        stats, warn, approx = build_slide(prs, svg)
        allres.append((stats, warn, approx))
        if not quiet:
            print('slide %d (%s): %s' % (i, Path(svg).name, ', '.join('%d %s' % (v, k) for k, v in stats.items())))
            for k, v in warn.items():
                print('  WARNING dropped: %s x%d' % (k, v))
            for k, v in approx.items():
                print('  note approximated: %s x%d' % (k, v))
    prs.save(out)
    if not quiet:
        print('wrote', out, len(svgs), 'slides (native objects)')
    return allres


def main(out, svgs):
    return convert(out, svgs)


if __name__ == '__main__':
    args = sys.argv[1:]
    if len(args) < 2:
        sys.exit(__doc__)
    main(args[0], args[1:])
