"""Export a rendered figure to the formats that journals and slides ask for.

python tools/export.py figure/out/<name> [--formats tif,eps,pptx,pptx-native]    (default: tif, eps, pptx)

<name> is the path of a figure that tools/render.mjs rendered, without the extension. The script reads <name>.png,
<name>.pdf and <name>.svg and writes next to them:
  <name>.tif   TIFF at the PNG's resolution, LZW compression (the raster file that many journals ask for)
  <name>.eps   EPS from the vector PDF (needs pdftops from poppler, or Ghostscript)
  <name>.pptx  one slide with the figure fitted inside 16:9 (needs python-pptx)
  <name>-native.pptx  (only with --formats pptx-native) the same slide made of real text boxes, shapes and groups
A missing tool is reported and skipped. Each file is opened again to check it, and the script exits with 1 when a
check fails.
"""
import re
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image

HERE = Path(__file__).resolve().parent


def tif(base):
    im = Image.open(f"{base}.png")
    dpi = tuple(round(v) for v in im.info.get("dpi", (600, 600)))
    if im.mode == "RGBA":   # an opaque page has no use for an alpha channel in a journal file
        flat = Image.new("RGB", im.size, "white")
        flat.paste(im, mask=im.split()[3])
        im = flat
    out = f"{base}.tif"
    im.convert("RGB").save(out, compression="tiff_lzw", dpi=dpi)
    back = Image.open(out)
    return out, f"{back.size[0]}x{back.size[1]} px, {round(back.info.get('dpi', (0,))[0])} dpi"


def eps(base):
    out = f"{base}.eps"
    if shutil.which("pdftops"):
        subprocess.run(["pdftops", "-eps", f"{base}.pdf", out], check=True)
    elif shutil.which("gs"):
        subprocess.run(["gs", "-q", "-dNOPAUSE", "-dBATCH", "-sDEVICE=eps2write", f"-sOutputFile={out}", f"{base}.pdf"], check=True)
    else:
        return None, "skipped: install poppler (pdftops) or Ghostscript"
    head = open(out, "rb").read(4096).decode("latin-1")
    box = re.search(r"%%BoundingBox:\s*(\d+) (\d+) (\d+) (\d+)", head)
    if not head.startswith("%!PS-Adobe") or "EPSF" not in head.split("\n")[0] or not box:
        raise RuntimeError("the EPS has no valid header or bounding box")
    return out, f"bounding box {box.group(3)}x{box.group(4)} pt"


def pptx(base):
    try:
        from pptx import Presentation
    except ImportError:
        return None, "skipped: pip install python-pptx"
    out = f"{base}.pptx"
    subprocess.run([sys.executable, str(HERE / "svg2pptx.py"), out, f"{base}.svg", f"{base}.png"], check=True, capture_output=True)
    p = Presentation(out)
    pic = p.slides[0].shapes[0]
    return out, f"{len(p.slides)} slide, picture {pic.width / pic.height:.2f}:1, slide {p.slide_width / p.slide_height:.2f}:1"


def pptx_native(base):
    """One slide of native PowerPoint objects (live text, shapes, groups). Only when asked for with --formats."""
    try:
        import pptx  # noqa: F401
    except ImportError:
        return None, "skipped: pip install python-pptx"
    out = f"{base}-native.pptx"
    r = subprocess.run([sys.executable, str(HERE / "svg2pptx.py"), "--native", out, f"{base}.svg", f"{base}.png"],
                       check=True, capture_output=True, text=True)
    from pptx import Presentation
    p = Presentation(out)
    n_text = n_shape = 0
    stack = list(p.slides[0].shapes)
    while stack:
        s = stack.pop()
        if s.shape_type == 6:   # group
            stack.extend(s.shapes)
        elif s.has_text_frame and s.text_frame.text.strip():
            n_text += 1
        else:
            n_shape += 1
    dropped = sum(1 for line in r.stdout.splitlines() if "WARNING" in line)
    return out, f"{len(p.slides)} slide, {n_text} text boxes, {n_shape} shapes, {dropped} kinds of dropped parts (see the converter output)"


FORMATS = {"tif": tif, "eps": eps, "pptx": pptx, "pptx-native": pptx_native}
DEFAULT_FORMATS = ["tif", "eps", "pptx"]   # pptx-native is added by --formats only

if __name__ == "__main__":
    args = sys.argv[1:]
    want = list(DEFAULT_FORMATS)
    if "--formats" in args:
        i = args.index("--formats")
        want = args[i + 1].split(",")
        del args[i:i + 2]
    if len(args) != 1 or any(w not in FORMATS for w in want):
        sys.exit(__doc__)
    base = args[0]
    for ext in ("png", "pdf", "svg"):
        if not Path(f"{base}.{ext}").exists():
            sys.exit(f"missing {base}.{ext}: render the figure with tools/render.mjs first")
    failed = 0
    for w in want:
        try:
            path, note = FORMATS[w](base)
            print(f"{w:<5} {Path(path).name + '  ' if path else ''}{note}")
        except Exception as e:   # report every format, then exit with 1
            failed += 1
            print(f"{w:<5} FAILED: {e}")
    sys.exit(1 if failed else 0)
